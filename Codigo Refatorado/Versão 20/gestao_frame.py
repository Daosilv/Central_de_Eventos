import tkinter as tk
from tkinter import ttk, messagebox, simpledialog, filedialog
import json
import openpyxl
import os
from PIL import Image, ImageTk

FONTE_PADRAO = ("Times New Roman", 12)
CAMINHO_LOGO = r'C:\BD_APP_Cemig\plemt_logo_ao_lado.PNG'

# --- DIÁLOGOS (sem alterações) ---
class MapEditDialog(tk.Toplevel):
    def __init__(self, parent, title, initial_data):
        super().__init__(parent)
        self.title(title)
        self.result = None
        self.transient(parent)
        self.grab_set()
        self.campo_var = tk.StringVar(value=initial_data.get('campo', ''))
        self.aba_var = tk.StringVar(value=initial_data.get('aba', ''))
        self.celula_var = tk.StringVar(value=initial_data.get('celula', ''))
        frame = tk.Frame(self, padx=15, pady=15)
        frame.pack(fill="both", expand=True)
        tk.Label(frame, text="Campo no Aplicativo:").grid(row=0, column=0, sticky='w', pady=2)
        tk.Entry(frame, textvariable=self.campo_var, width=40).grid(row=0, column=1, sticky='ew')
        tk.Label(frame, text="Aba no Excel:").grid(row=1, column=0, sticky='w', pady=2)
        tk.Entry(frame, textvariable=self.aba_var, width=40).grid(row=1, column=1, sticky='ew')
        tk.Label(frame, text="Célula no Excel:").grid(row=2, column=0, sticky='w', pady=2)
        tk.Entry(frame, textvariable=self.celula_var, width=40).grid(row=2, column=1, sticky='ew')
        button_frame = tk.Frame(frame)
        button_frame.grid(row=3, column=0, columnspan=2, pady=(10,0))
        ttk.Button(button_frame, text="Salvar", command=self.on_save).pack(side='left', padx=5)
        ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side='left', padx=5)
        self.wait_window(self)
    def on_save(self):
        self.result = {"campo_app": self.campo_var.get().strip(), "aba_excel": self.aba_var.get().strip(), "celula_excel": self.celula_var.get().strip()}
        self.destroy()

class UserDialog(tk.Toplevel):
    def __init__(self, parent, title, initial_data=None):
        super().__init__(parent)
        self.title(title); self.result = None; self.transient(parent); self.grab_set()
        self.nome_var = tk.StringVar(); self.matricula_var = tk.StringVar(); self.email_var = tk.StringVar()
        self.permissoes_vars = {"Visualização": tk.BooleanVar(), "Memorial de Cálculo": tk.BooleanVar(), "Ajustes": tk.BooleanVar(), "Gestão": tk.BooleanVar(), "Administrador": tk.BooleanVar()}
        self.perm_checkboxes = {}; self.dynamic_fields = {}
        main_frame = tk.Frame(self, padx=15, pady=15); main_frame.pack(fill="both", expand=True)
        info_frame = tk.LabelFrame(main_frame, text="Informações do Usuário", padx=10, pady=10); info_frame.pack(fill="x", expand=True)
        tk.Label(info_frame, text="Nome Completo:").grid(row=0, column=0, sticky='w', pady=2); tk.Entry(info_frame, textvariable=self.nome_var, width=40).grid(row=0, column=1, sticky='ew')
        tk.Label(info_frame, text="Matrícula:").grid(row=1, column=0, sticky='w', pady=2); tk.Entry(info_frame, textvariable=self.matricula_var, width=40).grid(row=1, column=1, sticky='ew')
        tk.Label(info_frame, text="Email:").grid(row=2, column=0, sticky='w', pady=2); tk.Entry(info_frame, textvariable=self.email_var, width=40).grid(row=2, column=1, sticky='ew')
        info_frame.grid_columnconfigure(1, weight=1)
        perm_frame = tk.LabelFrame(main_frame, text="Permissões", padx=10, pady=10); perm_frame.pack(fill="x", expand=True, pady=(10,0))
        for i, (nome_perm, var) in enumerate(self.permissoes_vars.items()):
            cb = ttk.Checkbutton(perm_frame, text=nome_perm, variable=var)
            if nome_perm == "Administrador": cb.config(command=self._on_admin_toggle)
            cb.grid(row=0, column=i, sticky='w', padx=5); self.perm_checkboxes[nome_perm] = cb
        self.dynamic_frame = tk.LabelFrame(main_frame, text="Informações Adicionais", padx=10, pady=10); self.dynamic_frame.pack(fill="x", expand=True, pady=(10,0)); self.dynamic_frame.grid_columnconfigure(1, weight=1)
        ttk.Button(self.dynamic_frame, text="Adicionar Campo", command=self.add_dynamic_field).grid(row=0, column=2, padx=5)
        if initial_data: self.populate_form(initial_data)
        button_frame = tk.Frame(main_frame); button_frame.pack(pady=(15,0)); ttk.Button(button_frame, text="Salvar", command=self.on_save).pack(side='left', padx=5); ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side='left', padx=5)
        self.wait_window(self)
    def populate_form(self, data):
        self.nome_var.set(data.get('nome_completo', '')); self.matricula_var.set(data.get('matricula', '')); self.email_var.set(data.get('email', '')); permissoes_str = data.get('permissoes', '')
        for nome, var in self.permissoes_vars.items():
            if nome in permissoes_str: var.set(True)
        dados_adicionais_json = data.get('dados_adicionais', '{}')
        if dados_adicionais_json:
            try:
                dados_adicionais = json.loads(dados_adicionais_json)
                for key, value in dados_adicionais.items(): self.add_dynamic_field(key, value)
            except json.JSONDecodeError: pass
        self._on_admin_toggle()
    def add_dynamic_field(self, field_name=None, field_value=""):
        if field_name is None: field_name = simpledialog.askstring("Novo Campo", "Digite o nome do novo campo (ex: CPF):", parent=self)
        if not field_name or not field_name.strip(): return
        field_name = field_name.strip()
        if field_name in self.dynamic_fields: messagebox.showwarning("Campo Duplicado", f"O campo '{field_name}' já existe.", parent=self); return
        row = len(self.dynamic_fields) + 1; var = tk.StringVar(value=field_value); self.dynamic_fields[field_name] = var
        tk.Label(self.dynamic_frame, text=f"{field_name}:").grid(row=row, column=0, sticky='w', pady=2); tk.Entry(self.dynamic_frame, textvariable=var, width=40).grid(row=row, column=1, sticky='ew')
    def _on_admin_toggle(self):
        is_admin = self.permissoes_vars["Administrador"].get(); new_state = tk.DISABLED if is_admin else tk.NORMAL
        for name, cb in self.perm_checkboxes.items():
            if name != "Administrador":
                cb.config(state=new_state)
                if is_admin: self.permissoes_vars[name].set(True)
    def on_save(self):
        nome = self.nome_var.get().strip(); matricula = self.matricula_var.get().strip()
        if not nome or not matricula: messagebox.showerror("Erro", "Os campos 'Nome Completo' e 'Matrícula' são obrigatórios.", parent=self); return
        permissoes_list = [nome for nome, var in self.permissoes_vars.items() if var.get()]; permissoes_str = ",".join(permissoes_list); dynamic_data = {key: var.get() for key, var in self.dynamic_fields.items()}
        self.result = {"nome_completo": nome, "matricula": matricula, "email": self.email_var.get().strip(), "permissoes": permissoes_str, "dados_adicionais": json.dumps(dynamic_data)}; self.destroy()

class MultiAddDialog(tk.Toplevel):
    def __init__(self, parent, title):
        super().__init__(parent); self.result = None; self.title(title); self.geometry("400x400"); self.transient(parent); self.grab_set()
        main_frame = tk.Frame(self, padx=10, pady=10); main_frame.pack(fill="both", expand=True)
        tk.Label(main_frame, text="Cole a lista de itens abaixo (um por linha):", font=FONTE_PADRAO).pack(side="top", anchor="w")
        button_frame = tk.Frame(main_frame); button_frame.pack(side="bottom", fill="x", pady=(5,0))
        text_frame = tk.Frame(main_frame, bd=1, relief="solid"); text_frame.pack(side="top", fill="both", expand=True, pady=5)
        self.text_widget = tk.Text(text_frame, wrap="word", font=FONTE_PADRAO); self.text_widget.pack(fill="both", expand=True, padx=2, pady=2)
        ttk.Button(button_frame, text="Adicionar", command=self.on_add).pack(side="right", padx=5); ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side="right"); self.wait_window(self)
    def on_add(self):
        text = self.text_widget.get("1.0", "end-1c"); self.result = [line.strip() for line in text.splitlines() if line.strip()]
        if not self.result: messagebox.showwarning("Nenhum item", "A lista está vazia.", parent=self); return
        self.destroy()

class MultiCaboDialog(tk.Toplevel):
    def __init__(self, parent, title):
        super().__init__(parent); self.result = None; self.title(title); self.geometry("500x400"); self.transient(parent); self.grab_set()
        main_frame = tk.Frame(self, padx=10, pady=10); main_frame.pack(fill="both", expand=True)
        instructions = "Cole os dados dos cabos abaixo usando o formato:\nTrecho,Nominal,Admissível (um por linha)"; tk.Label(main_frame, text=instructions, font=FONTE_PADRAO, justify='left').pack(side="top", anchor="w")
        button_frame = tk.Frame(main_frame); button_frame.pack(side="bottom", fill="x", pady=(5,0))
        text_frame = tk.Frame(main_frame, bd=1, relief="solid"); text_frame.pack(side="top", fill="both", expand=True, pady=5)
        self.text_widget = tk.Text(text_frame, wrap="word", font=FONTE_PADRAO); self.text_widget.pack(fill="both", expand=True, padx=2, pady=2)
        ttk.Button(button_frame, text="Adicionar", command=self.on_add).pack(side="right", padx=5); ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side="right"); self.wait_window(self)
    def on_add(self):
        text = self.text_widget.get("1.0", "end-1c"); lines = [line.strip() for line in text.splitlines() if line.strip()]; parsed_data = []
        for i, line in enumerate(lines, 1):
            parts = [p.strip() for p in line.split(',')]
            if len(parts) != 3: messagebox.showerror("Formato Inválido", f"A linha {i} está com formato incorreto.\nUse: Trecho,Nominal,Admissível", parent=self); return
            parsed_data.append(parts)
        if not parsed_data: messagebox.showwarning("Nenhum item", "A lista está vazia ou em formato inválido.", parent=self); return
        self.result = parsed_data; self.destroy()

class MultiMapDialog(tk.Toplevel):
    def __init__(self, parent, title):
        super().__init__(parent); self.result = None; self.title(title); self.geometry("500x400"); self.transient(parent); self.grab_set()
        main_frame = tk.Frame(self, padx=10, pady=10); main_frame.pack(fill="both", expand=True)
        instructions = "Cole os dados do mapeamento abaixo usando o formato:\nCampo_App,Aba_Excel,Celula_Excel (um por linha)"; tk.Label(main_frame, text=instructions, font=FONTE_PADRAO, justify='left').pack(side="top", anchor="w")
        button_frame = tk.Frame(main_frame); button_frame.pack(side="bottom", fill="x", pady=(5,0))
        text_frame = tk.Frame(main_frame, bd=1, relief="solid"); text_frame.pack(side="top", fill="both", expand=True, pady=5)
        self.text_widget = tk.Text(text_frame, wrap="word", font=FONTE_PADRAO); self.text_widget.pack(fill="both", expand=True, padx=2, pady=2)
        ttk.Button(button_frame, text="Adicionar", command=self.on_add).pack(side="right", padx=5); ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side="right"); self.wait_window(self)
    def on_add(self):
        text = self.text_widget.get("1.0", "end-1c"); lines = [line.strip() for line in text.splitlines() if line.strip()]; parsed_data = []
        for i, line in enumerate(lines, 1):
            parts = [p.strip() for p in line.split(',')]
            if len(parts) != 3: messagebox.showerror("Formato Inválido", f"A linha {i} está com formato incorreto.\nUse: Campo,Aba,Célula", parent=self); return
            parsed_data.append(parts)
        if not parsed_data: messagebox.showwarning("Nenhum item", "A lista está vazia ou em formato inválido.", parent=self); return
        self.result = parsed_data; self.destroy()

class CaboDialog(tk.Toplevel):
    def __init__(self, parent, title, initial_data=None):
        super().__init__(parent); self.title(title); self.result = None; self.transient(parent); self.grab_set()
        frame = tk.Frame(self, padx=15, pady=15); frame.pack(fill="both", expand=True)
        tk.Label(frame, text="Trecho:").grid(row=0, column=0, sticky='w', pady=2); self.trecho_var = tk.StringVar(); tk.Entry(frame, textvariable=self.trecho_var, width=30).grid(row=0, column=1, sticky='ew')
        tk.Label(frame, text="Nominal (A):").grid(row=1, column=0, sticky='w', pady=2); self.nominal_var = tk.StringVar(); tk.Entry(frame, textvariable=self.nominal_var, width=30).grid(row=1, column=1, sticky='ew')
        tk.Label(frame, text="Admissível (A):").grid(row=2, column=0, sticky='w', pady=2); self.admissivel_var = tk.StringVar(); tk.Entry(frame, textvariable=self.admissivel_var, width=30).grid(row=2, column=1, sticky='ew')
        if initial_data: self.trecho_var.set(initial_data[0]); self.nominal_var.set(initial_data[1]); self.admissivel_var.set(initial_data[2])
        button_frame = tk.Frame(frame); button_frame.grid(row=3, column=0, columnspan=2, pady=(10,0)); ttk.Button(button_frame, text="Salvar", command=self.on_save).pack(side='left', padx=5); ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side='left', padx=5)
        self.wait_window(self)
    def on_save(self):
        trecho = self.trecho_var.get().strip(); nominal = self.nominal_var.get().strip(); admissivel = self.admissivel_var.get().strip()
        if not trecho: messagebox.showerror("Erro", "O campo 'Trecho' é obrigatório.", parent=self); return
        self.result = (trecho, nominal, admissivel); self.destroy()

# --- NOVO DIÁLOGO PARA CONFIGURAÇÕES ---
class ConfigDialog(tk.Toplevel):
    def __init__(self, parent, title, num_fields, initial_data=None):
        super().__init__(parent)
        self.title(title)
        self.result = None
        self.num_fields = num_fields
        self.transient(parent)
        self.grab_set()

        self.nome_var = tk.StringVar()
        self.value_vars = [tk.StringVar() for _ in range(num_fields)]

        frame = tk.Frame(self, padx=15, pady=15)
        frame.pack(fill="both", expand=True)

        tk.Label(frame, text="Nome da Configuração:").grid(row=0, column=0, sticky='w', pady=2)
        self.nome_entry = tk.Entry(frame, textvariable=self.nome_var, width=40)
        self.nome_entry.grid(row=0, column=1, sticky='ew')
        
        for i in range(num_fields):
            tk.Label(frame, text=f"Valor {i+1}:").grid(row=i+1, column=0, sticky='w', pady=2)
            tk.Entry(frame, textvariable=self.value_vars[i], width=40).grid(row=i+1, column=1, sticky='ew')

        if initial_data:
            self.nome_var.set(initial_data.get('nome', ''))
            for i in range(num_fields):
                self.value_vars[i].set(initial_data.get(f'v{i+1}', ''))
            if initial_data.get('nome', '').endswith('Padrão'):
                self.nome_entry.config(state='readonly')

        button_frame = tk.Frame(frame)
        button_frame.grid(row=num_fields + 1, column=0, columnspan=2, pady=(10, 0))
        ttk.Button(button_frame, text="Salvar", command=self.on_save).pack(side='left', padx=5)
        ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side='left', padx=5)
        self.wait_window(self)

    def on_save(self):
        nome = self.nome_var.get().strip()
        if not nome:
            messagebox.showerror("Erro", "O campo 'Nome da Configuração' é obrigatório.", parent=self)
            return
        
        self.result = {'nome': nome}
        for i in range(self.num_fields):
            self.result[f'v{i+1}'] = self.value_vars[i].get().strip()
        self.destroy()

# --- NOVO DIÁLOGO PARA A ABA AJUDA ---
class AjudaEditDialog(tk.Toplevel):
    def __init__(self, parent, controller, ajuda_id):
        super().__init__(parent)
        self.controller = controller
        self.ajuda_id = ajuda_id
        
        self.initial_data = self.controller.get_ajuda_by_id(self.ajuda_id)
        if not self.initial_data:
            messagebox.showerror("Erro", "Não foi possível carregar os dados da ajuda.", parent=parent)
            self.destroy()
            return

        self.title("Editar Ajuda")
        self.geometry("700x500")
        self.transient(parent)
        self.grab_set()

        self.titulo_var = tk.StringVar(value=self.initial_data.get('titulo', ''))
        self.anexos_list = json.loads(self.initial_data.get('anexos', '[]'))
        
        main_frame = tk.Frame(self, padx=15, pady=15)
        main_frame.pack(fill='both', expand=True)
        
        # Título
        title_frame = tk.Frame(main_frame)
        title_frame.pack(fill='x', pady=(0, 10))
        tk.Label(title_frame, text="Título:", font=FONTE_PADRAO).pack(side='left', padx=(0, 5))
        tk.Entry(title_frame, textvariable=self.titulo_var, font=FONTE_PADRAO).pack(side='left', fill='x', expand=True)

        # Anexos
        anexos_frame = tk.LabelFrame(main_frame, text="Anexos", padx=10, pady=10)
        anexos_frame.pack(fill='x', pady=(0, 10))
        self.anexos_listbox = tk.Listbox(anexos_frame, height=4, font=("Times New Roman", 10))
        self.anexos_listbox.pack(side='left', fill='x', expand=True, padx=(0,10))
        anexos_btn_frame = tk.Frame(anexos_frame)
        anexos_btn_frame.pack(side='left')
        ttk.Button(anexos_btn_frame, text="Adicionar", command=self.adicionar_anexo).pack(fill='x')
        ttk.Button(anexos_btn_frame, text="Remover", command=self.remover_anexo).pack(fill='x', pady=5)
        self.carregar_anexos_listbox()

        # Texto
        text_frame = tk.LabelFrame(main_frame, text="Conteúdo", padx=10, pady=10)
        text_frame.pack(fill='both', expand=True)
        self.text_widget = tk.Text(text_frame, wrap='word', font=FONTE_PADRAO)
        self.text_widget.pack(fill='both', expand=True)
        self.text_widget.insert('1.0', self.initial_data.get('texto', ''))

        # Botões Salvar/Cancelar
        button_frame = tk.Frame(main_frame)
        button_frame.pack(pady=(10,0))
        ttk.Button(button_frame, text="Salvar", command=self.on_save).pack(side='left', padx=5)
        ttk.Button(button_frame, text="Cancelar", command=self.destroy).pack(side='left', padx=5)
        
        self.wait_window(self)

    def carregar_anexos_listbox(self):
        self.anexos_listbox.delete(0, tk.END)
        for anexo in self.anexos_list:
            self.anexos_listbox.insert(tk.END, os.path.basename(anexo))

    def adicionar_anexo(self):
        paths = filedialog.askopenfilenames(title="Selecionar arquivos", parent=self)
        if paths:
            for p in paths:
                if p not in self.anexos_list:
                    self.anexos_list.append(p)
            self.carregar_anexos_listbox()
    
    def remover_anexo(self):
        indices = self.anexos_listbox.curselection()
        if not indices: return
        for i in sorted(indices, reverse=True):
            self.anexos_list.pop(i)
        self.carregar_anexos_listbox()

    def on_save(self):
        titulo = self.titulo_var.get().strip()
        texto = self.text_widget.get('1.0', 'end-1c').strip()
        anexos_json = json.dumps(self.anexos_list)
        if not titulo:
            messagebox.showerror("Erro", "O campo 'Título' é obrigatório.", parent=self)
            return
        
        self.controller.update_ajuda(self.ajuda_id, titulo, texto, anexos_json)
        self.destroy()

# --- FRAME PRINCIPAL ---
class GestaoFrame(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        
        self._criar_logo_header(self)

        notebook = ttk.Notebook(self)
        notebook.pack(expand=True, fill='both', padx=10, pady=10)

        tab1 = TabListasGerais(notebook, controller)
        tab2 = TabHierarquia(notebook, controller)
        tab3 = TabCabos(notebook, controller)
        tab4 = TabMapeamento(notebook, controller)
        tab5 = TabUsuario(notebook, controller)
        tab6 = TabAjuda(notebook, controller) # NOVA ABA

        notebook.add(tab1, text='Listas Gerais')
        notebook.add(tab2, text='Hierarquia')
        notebook.add(tab3, text='Cabos Críticos')
        notebook.add(tab4, text='Mapeamento Excel')
        notebook.add(tab5, text='Usuários')
        notebook.add(tab6, text='Ajuda') # NOVA ABA

    def _criar_logo_header(self, parent):
        logo_frame = tk.Frame(parent)
        logo_frame.pack(fill='x', pady=10)
        try:
            if os.path.exists(CAMINHO_LOGO):
                img = Image.open(CAMINHO_LOGO)
                
                max_height = 80 
                img_ratio = img.width / img.height
                new_height = max_height
                new_width = int(new_height * img_ratio)
                
                img = img.resize((new_width, new_height), Image.Resampling.LANCZOS)
                
                self.logo_photo = ImageTk.PhotoImage(img)
                logo_label = tk.Label(logo_frame, image=self.logo_photo)
                logo_label.pack()
            else:
                tk.Label(logo_frame, text="Plemt", font=("Helvetica", 24, "bold")).pack()
        except Exception as e:
            print(f"Erro ao carregar o logo: {e}")
            tk.Label(logo_frame, text="Plemt", font=("Helvetica", 24, "bold")).pack()


# --- ABAS MODIFICADAS E NOVAS ---

class TabUsuario(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent); self.controller = controller; self.users_data = {}; self.criar_widgets(); self.carregar_usuarios()
    def criar_widgets(self):
        main_frame = tk.Frame(self, padx=10, pady=10); main_frame.pack(fill="both", expand=True); main_frame.grid_columnconfigure(0, weight=1); main_frame.grid_rowconfigure(0, weight=1)
        tree_frame = tk.LabelFrame(main_frame, text="Usuários Cadastrados"); tree_frame.grid(row=0, column=0, sticky='nsew', padx=(0, 10)); tree_frame.grid_rowconfigure(0, weight=1); tree_frame.grid_columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(tree_frame, columns=('nome', 'matricula', 'email', 'permissoes'), show='headings'); self.tree.heading('nome', text='Nome Completo'); self.tree.heading('matricula', text='Matrícula'); self.tree.heading('email', text='Email'); self.tree.heading('permissoes', text='Permissões')
        self.tree.column('nome', width=250); self.tree.column('matricula', width=100); self.tree.column('email', width=200); self.tree.column('permissoes', width=300); self.tree.pack(fill='both', expand=True)
        botoes_frame = tk.Frame(main_frame); botoes_frame.grid(row=0, column=1, sticky="ns"); ttk.Button(botoes_frame, text="Adicionar Usuário", command=self.adicionar_usuario).pack(pady=5, fill='x'); ttk.Button(botoes_frame, text="Editar Usuário", command=self.editar_usuario).pack(pady=5, fill='x'); ttk.Button(botoes_frame, text="Remover Usuário", command=self.remover_usuario).pack(pady=5, fill='x')
    def carregar_usuarios(self):
        self.tree.delete(*self.tree.get_children()); self.users_data.clear(); usuarios = self.controller.get_all_users()
        if not usuarios: return
        for user in usuarios:
            user_id = user['id']; self.users_data[user_id] = user
            self.tree.insert('', 'end', iid=user_id, values=(user['nome_completo'], user['matricula'], user['email'], user['permissoes']))
    def adicionar_usuario(self):
        dialog = UserDialog(self, title="Adicionar Novo Usuário")
        if dialog.result: self.controller.add_user(dialog.result); self.carregar_usuarios()
    def editar_usuario(self):
        selected_id = self.tree.focus()
        if not selected_id: messagebox.showwarning("Atenção", "Selecione um usuário para editar.", parent=self); return
        user_data = self.users_data.get(int(selected_id))
        if not user_data: return
        dialog = UserDialog(self, title="Editar Usuário", initial_data=user_data)
        if dialog.result: self.controller.update_user(int(selected_id), dialog.result); self.carregar_usuarios()
    def remover_usuario(self):
        selected_id = self.tree.focus()
        if not selected_id: messagebox.showwarning("Atenção", "Selecione um usuário para remover.", parent=self); return
        if messagebox.askyesno("Confirmar", "Deseja remover o usuário selecionado?", parent=self): self.controller.delete_user(int(selected_id)); self.carregar_usuarios()

class TabListasGerais(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        self.criar_widgets()
        self.carregar_categorias()

    def criar_widgets(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(1, weight=1)
        
        frame_superior = tk.LabelFrame(self, text="Gerenciar Opções das Listas", padx=10, pady=10)
        frame_superior.grid(row=0, column=0, padx=10, pady=10, sticky="ew")
        frame_superior.grid_columnconfigure(1, weight=1)
        tk.Label(frame_superior, text="Selecione a Categoria:", font=FONTE_PADRAO).grid(row=0, column=0, padx=(0, 5), sticky='w')
        self.combo_categorias = ttk.Combobox(frame_superior, state="readonly", font=FONTE_PADRAO)
        self.combo_categorias.grid(row=0, column=1, sticky="ew")
        self.combo_categorias.bind("<<ComboboxSelected>>", self.on_categoria_selecionada)
        
        self.container_frame = tk.Frame(self)
        self.container_frame.grid(row=1, column=0, padx=10, pady=(0, 10), sticky="nsew")
        self.container_frame.grid_rowconfigure(0, weight=1)
        self.container_frame.grid_columnconfigure(0, weight=1)

        self._criar_frame_listas_simples()
        self._criar_frame_config("tempo_morto", 4)
        self._criar_frame_config("grupo4", 10)

    def _criar_frame_listas_simples(self):
        self.frame_listas_simples = tk.Frame(self.container_frame)
        self.frame_listas_simples.grid(row=0, column=0, sticky="nsew")
        self.frame_listas_simples.grid_rowconfigure(0, weight=1)
        self.frame_listas_simples.grid_columnconfigure(0, weight=1)
        self.id_map_simples = {}

        list_frame = tk.LabelFrame(self.frame_listas_simples, text="Opções", padx=5, pady=5)
        list_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        list_frame.grid_rowconfigure(0, weight=1); list_frame.grid_columnconfigure(0, weight=1)
        scrollbar = ttk.Scrollbar(list_frame); scrollbar.grid(row=0, column=1, sticky='ns')
        self.listbox_opcoes = tk.Listbox(list_frame, yscrollcommand=scrollbar.set, font=FONTE_PADRAO)
        self.listbox_opcoes.grid(row=0, column=0, sticky='nsew'); scrollbar.config(command=self.listbox_opcoes.yview)
        
        botoes_frame = tk.Frame(self.frame_listas_simples)
        botoes_frame.grid(row=0, column=1, sticky="ns")
        ttk.Button(botoes_frame, text="Adicionar", command=self.adicionar_opcao_simples).pack(pady=5, fill='x')
        ttk.Button(botoes_frame, text="Editar", command=self.editar_opcao_simples).pack(pady=5, fill='x')
        ttk.Button(botoes_frame, text="Remover", command=self.remover_opcao_simples).pack(pady=5, fill='x')
        ttk.Button(botoes_frame, text="Remover Tudo", command=self.remover_todas_opcoes_simples).pack(pady=(15, 5), fill='x')

    def _criar_frame_config(self, nome, num_valores):
        frame = tk.Frame(self.container_frame)
        frame.grid(row=0, column=0, sticky="nsew")
        frame.grid_rowconfigure(0, weight=1)
        frame.grid_columnconfigure(0, weight=1)
        
        tree_frame = tk.LabelFrame(frame, text="Configurações", padx=5, pady=5)
        tree_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        tree_frame.grid_rowconfigure(0, weight=1); tree_frame.grid_columnconfigure(0, weight=1)
        
        cols = ['nome'] + [f"v{i+1}" for i in range(num_valores)]
        headings = ['Nome'] + [f"Valor {i+1}" for i in range(num_valores)]
        tree = ttk.Treeview(tree_frame, columns=cols, show='headings')
        for col, heading in zip(cols, headings):
            tree.heading(col, text=heading)
            tree.column(col, width=100, anchor='center')
        tree.pack(fill='both', expand=True)
        
        botoes_frame = tk.Frame(frame)
        botoes_frame.grid(row=0, column=1, sticky="ns")
        ttk.Button(botoes_frame, text="Adicionar", command=lambda: self.adicionar_config(nome, num_valores)).pack(pady=5, fill='x')
        ttk.Button(botoes_frame, text="Editar", command=lambda: self.editar_config(nome, num_valores)).pack(pady=5, fill='x')
        ttk.Button(botoes_frame, text="Remover", command=lambda: self.remover_config(nome)).pack(pady=5, fill='x')
        
        setattr(self, f"frame_{nome}", frame)
        setattr(self, f"tree_{nome}", tree)

    def carregar_categorias(self):
        categorias = self.controller.get_categorias_gerais() + ["Tempo Morto", "Grupo 4"]
        self.combo_categorias['values'] = sorted(categorias)
        if self.combo_categorias['values']:
            self.combo_categorias.set(self.combo_categorias['values'][0])
            self.on_categoria_selecionada()

    def on_categoria_selecionada(self, event=None):
        categoria = self.combo_categorias.get()
        if not categoria: return

        if categoria == "Tempo Morto":
            self.frame_tempo_morto.tkraise()
            self.carregar_dados_config("tempo_morto")
        elif categoria == "Grupo 4":
            self.frame_grupo4.tkraise()
            self.carregar_dados_config("grupo4")
        else:
            self.frame_listas_simples.tkraise()
            self.carregar_opcoes_simples()
    
    # --- Métodos para Listas Simples ---
    def carregar_opcoes_simples(self):
        categoria = self.combo_categorias.get()
        self.listbox_opcoes.delete(0, tk.END); self.id_map_simples.clear()
        opcoes = self.controller.get_opcoes_por_categoria_geral(categoria)
        for item in opcoes:
            self.listbox_opcoes.insert(tk.END, item['valor'])
            self.id_map_simples[self.listbox_opcoes.size() - 1] = item['id']

    def adicionar_opcao_simples(self):
        categoria = self.combo_categorias.get()
        dialog = MultiAddDialog(self, f"Adicionar para '{categoria}'")
        if dialog.result:
            for valor in dialog.result: self.controller.adicionar_opcao_geral(categoria, valor)
            self.carregar_opcoes_simples()
            messagebox.showinfo("Sucesso", f"{len(dialog.result)} item(ns) adicionado(s).", parent=self)
    
    def editar_opcao_simples(self):
        idx = self.listbox_opcoes.curselection()
        if not idx: messagebox.showwarning("Atenção", "Selecione uma opção.", parent=self); return
        id_opcao = self.id_map_simples[idx[0]]; valor_atual = self.listbox_opcoes.get(idx[0])
        novo_valor = simpledialog.askstring("Editar", "Novo valor:", initialvalue=valor_atual, parent=self)
        if novo_valor and novo_valor.strip():
            self.controller.editar_opcao_geral(id_opcao, novo_valor.strip())
            self.carregar_opcoes_simples()
            
    def remover_opcao_simples(self):
        idx = self.listbox_opcoes.curselection()
        if not idx: messagebox.showwarning("Atenção", "Selecione uma opção.", parent=self); return
        id_opcao = self.id_map_simples[idx[0]]
        if messagebox.askyesno("Confirmar", "Deseja remover a opção selecionada?", parent=self):
            self.controller.deletar_opcao_geral(id_opcao)
            self.carregar_opcoes_simples()

    def remover_todas_opcoes_simples(self):
        categoria = self.combo_categorias.get()
        if messagebox.askyesno("Confirmar", f"Deseja remover TODAS as opções de '{categoria}'?", icon='warning', parent=self):
            self.controller.deletar_opcoes_por_categoria_geral(categoria)
            self.carregar_opcoes_simples()
            
    # --- Métodos Genéricos para Configurações (Tempo Morto, Grupo 4) ---
    def carregar_dados_config(self, nome):
        tree = getattr(self, f"tree_{nome}")
        tree.delete(*tree.get_children())
        
        get_all_func = getattr(self.controller, f"get_all_{nome}")
        configs = get_all_func()
        if not configs: return
        
        for item in configs:
            values = [item.get(key, '') for key in tree['columns']]
            tree.insert('', 'end', iid=item['id'], values=values)

    def adicionar_config(self, nome, num_valores):
        dialog = ConfigDialog(self, f"Adicionar {nome.replace('_', ' ').title()}", num_valores)
        if dialog.result:
            add_func = getattr(self.controller, f"add_{nome}")
            add_func(dialog.result)
            self.carregar_dados_config(nome)
            
    def editar_config(self, nome, num_valores):
        tree = getattr(self, f"tree_{nome}")
        selected_id = tree.focus()
        if not selected_id: messagebox.showwarning("Atenção", "Selecione um item para editar.", parent=self); return
        
        item_data = tree.item(selected_id, 'values')
        initial_data = {'nome': item_data[0]}
        for i in range(num_valores):
            initial_data[f'v{i+1}'] = item_data[i+1]
            
        dialog = ConfigDialog(self, f"Editar {nome.replace('_', ' ').title()}", num_valores, initial_data)
        if dialog.result:
            update_func = getattr(self.controller, f"update_{nome}")
            update_func(int(selected_id), dialog.result)
            self.carregar_dados_config(nome)

    def remover_config(self, nome):
        tree = getattr(self, f"tree_{nome}")
        selected_id = tree.focus()
        if not selected_id: messagebox.showwarning("Atenção", "Selecione um item para remover.", parent=self); return
        
        item_data = tree.item(selected_id, 'values')
        if item_data[0].endswith('Padrão'):
            messagebox.showerror("Erro", "A configuração padrão não pode ser removida.", parent=self)
            return
            
        if messagebox.askyesno("Confirmar", "Deseja remover o item selecionado?", parent=self):
            delete_func = getattr(self.controller, f"delete_{nome}")
            delete_func(int(selected_id))
            self.carregar_dados_config(nome)

# --- ABA DE AJUDA ---
class TabAjuda(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent)
        self.controller = controller
        self.id_map = {}
        self.criar_widgets()
        self.carregar_ajuda()

    def criar_widgets(self):
        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)
        
        main_frame = tk.Frame(self, padx=10, pady=10)
        main_frame.pack(fill="both", expand=True)
        main_frame.grid_columnconfigure(0, weight=1)
        main_frame.grid_rowconfigure(0, weight=1)

        list_frame = tk.LabelFrame(main_frame, text="Tópicos de Ajuda", padx=5, pady=5)
        list_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        list_frame.grid_rowconfigure(0, weight=1)
        list_frame.grid_columnconfigure(0, weight=1)

        self.listbox_ajuda = tk.Listbox(list_frame, font=FONTE_PADRAO)
        self.listbox_ajuda.pack(fill='both', expand=True)
        self.listbox_ajuda.bind("<Double-1>", lambda e: self.editar_ajuda())

        botoes_frame = tk.Frame(main_frame)
        botoes_frame.grid(row=0, column=1, sticky="ns")
        ttk.Button(botoes_frame, text="Adicionar Linha", command=self.adicionar_ajuda).pack(pady=5, fill='x')
        ttk.Button(botoes_frame, text="Editar Linha", command=self.editar_ajuda).pack(pady=5, fill='x')
        ttk.Button(botoes_frame, text="Excluir Linha", command=self.excluir_ajuda).pack(pady=5, fill='x')

    def carregar_ajuda(self):
        self.listbox_ajuda.delete(0, tk.END)
        self.id_map.clear()
        ajuda_items = self.controller.get_all_ajuda()
        if not ajuda_items: return
        for item in ajuda_items:
            self.listbox_ajuda.insert(tk.END, item['titulo'])
            self.id_map[self.listbox_ajuda.size() - 1] = item['id']
            
    def adicionar_ajuda(self):
        titulo = simpledialog.askstring("Novo Tópico", "Digite o título para a nova linha de ajuda:", parent=self)
        if titulo and titulo.strip():
            self.controller.add_ajuda(titulo.strip())
            self.carregar_ajuda()

    def editar_ajuda(self):
        selected_idx = self.listbox_ajuda.curselection()
        if not selected_idx:
            messagebox.showwarning("Atenção", "Selecione um tópico para editar.", parent=self)
            return
        ajuda_id = self.id_map.get(selected_idx[0])
        if ajuda_id:
            dialog = AjudaEditDialog(self, self.controller, ajuda_id)
            self.carregar_ajuda()

    def excluir_ajuda(self):
        selected_idx = self.listbox_ajuda.curselection()
        if not selected_idx:
            messagebox.showwarning("Atenção", "Selecione um tópico para excluir.", parent=self)
            return
        
        ajuda_id = self.id_map.get(selected_idx[0])
        if messagebox.askyesno("Confirmar", "Deseja excluir o tópico de ajuda selecionado?", parent=self):
            self.controller.delete_ajuda(ajuda_id)
            self.carregar_ajuda()
            
# --- ABAS EXISTENTES (sem alterações) ---
class TabHierarquia(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent); self.controller = controller; self.criar_widgets(); self.carregar_hierarquia()
    def criar_widgets(self):
        main_frame = tk.Frame(self, padx=10, pady=10); main_frame.pack(fill="both", expand=True); main_frame.grid_columnconfigure(0, weight=1); main_frame.grid_rowconfigure(0, weight=1)
        tree_frame = tk.LabelFrame(main_frame, text="Estrutura"); tree_frame.grid(row=0, column=0, sticky='nsew', padx=(0, 10)); tree_frame.grid_rowconfigure(0, weight=1); tree_frame.grid_columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(tree_frame, columns=('valor'), show='tree headings'); self.tree.heading('#0', text='Item'); self.tree.heading('valor', text='Categoria'); self.tree.pack(fill='both', expand=True); self.tree.bind('<<TreeviewSelect>>', self.on_tree_select)
        botoes_frame = tk.Frame(main_frame); botoes_frame.grid(row=0, column=1, sticky="ns")
        self.btn_add_fab = ttk.Button(botoes_frame, text="Adicionar Fabricante", command=lambda: self.adicionar_item('fabricante')); self.btn_add_fab.pack(pady=5, fill='x')
        self.btn_add_mod = ttk.Button(botoes_frame, text="Adicionar Modelo", command=lambda: self.adicionar_item('modelo'), state='disabled'); self.btn_add_mod.pack(pady=5, fill='x')
        self.btn_add_com = ttk.Button(botoes_frame, text="Adicionar Comando", command=lambda: self.adicionar_item('comando'), state='disabled'); self.btn_add_com.pack(pady=5, fill='x')
        self.btn_add_multi = ttk.Button(botoes_frame, text="Adicionar Múltiplos", command=self.adicionar_multiplos_itens, state='disabled'); self.btn_add_multi.pack(pady=(5, 10), fill='x')
        ttk.Separator(botoes_frame, orient='horizontal').pack(pady=5, fill='x')
        self.btn_edit = ttk.Button(botoes_frame, text="Editar Selecionado", state='disabled', command=self.editar_item); self.btn_edit.pack(pady=5, fill='x')
        self.btn_del = ttk.Button(botoes_frame, text="Remover Selecionado", state='disabled', command=self.remover_item); self.btn_del.pack(pady=5, fill='x')
    def carregar_hierarquia(self, parent_node="", id_pai=None):
        if parent_node == "": self.tree.delete(*self.tree.get_children())
        items = self.controller.get_opcoes_hierarquicas(id_pai)
        for item in items:
            node_id = self.tree.insert(parent_node, 'end', iid=item['id'], text=item['valor'], values=(item['categoria'].title(),)); self.carregar_hierarquia(node_id, item['id'])
    def on_tree_select(self, event):
        selected_id = self.tree.focus()
        if not selected_id: self.btn_add_mod.config(state='disabled'); self.btn_add_com.config(state='disabled'); self.btn_add_multi.config(state='disabled'); self.btn_edit.config(state='disabled'); self.btn_del.config(state='disabled'); return
        self.btn_edit.config(state='normal'); self.btn_del.config(state='normal'); item = self.tree.item(selected_id); categoria = item['values'][0].lower()
        if categoria == 'fabricante': self.btn_add_mod.config(state='normal'); self.btn_add_com.config(state='disabled'); self.btn_add_multi.config(state='normal')
        elif categoria == 'modelo': self.btn_add_mod.config(state='disabled'); self.btn_add_com.config(state='normal'); self.btn_add_multi.config(state='normal')
        else: self.btn_add_mod.config(state='disabled'); self.btn_add_com.config(state='disabled'); self.btn_add_multi.config(state='disabled')
    def adicionar_item(self, categoria):
        id_pai = None; prompt = f"Novo valor para {categoria.title()}:"
        if categoria != 'fabricante':
            selected_id = self.tree.focus()
            if not selected_id: messagebox.showerror("Erro", "Nenhum item pai selecionado.", parent=self); return
            id_pai = int(selected_id)
        novo_valor = simpledialog.askstring("Adicionar Item", prompt, parent=self)
        if novo_valor and novo_valor.strip(): self.controller.adicionar_opcao_hierarquica(categoria, novo_valor.strip(), id_pai); self.carregar_hierarquia()
    def adicionar_multiplos_itens(self):
        selected_id = self.tree.focus();
        if not selected_id: return
        id_pai = int(selected_id); item = self.tree.item(selected_id); categoria_pai = item['values'][0].lower()
        if categoria_pai == 'fabricante': categoria_filho = 'modelo'
        elif categoria_pai == 'modelo': categoria_filho = 'comando'
        else: return
        dialog = MultiAddDialog(self, f"Adicionar Múltiplos para '{item['text']}'")
        if dialog.result:
            for valor in dialog.result: self.controller.adicionar_opcao_hierarquica(categoria_filho, valor, id_pai)
            self.carregar_hierarquia(); self.tree.selection_set(selected_id)
    def editar_item(self):
        selected_id = self.tree.focus();
        if not selected_id: return
        valor_atual = self.tree.item(selected_id, 'text'); novo_valor = simpledialog.askstring("Editar", "Novo valor:", initialvalue=valor_atual, parent=self)
        if novo_valor and novo_valor.strip(): self.controller.editar_opcao_hierarquica(int(selected_id), novo_valor.strip()); self.carregar_hierarquia()
    def remover_item(self):
        selected_id = self.tree.focus();
        if not selected_id: return
        if messagebox.askyesno("Confirmar", "Deseja remover o item e todos os seus filhos?", icon='warning', parent=self): self.controller.deletar_opcao_hierarquica(int(selected_id)); self.carregar_hierarquia()

class TabCabos(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent); self.controller = controller; self.criar_widgets(); self.carregar_cabos()
    def criar_widgets(self):
        main_frame = tk.Frame(self, padx=10, pady=10); main_frame.pack(fill="both", expand=True); main_frame.grid_columnconfigure(0, weight=1); main_frame.grid_rowconfigure(0, weight=1)
        tree_frame = tk.LabelFrame(main_frame, text="Propriedades dos Cabos Críticos"); tree_frame.grid(row=0, column=0, sticky='nsew', padx=(0, 10)); tree_frame.grid_rowconfigure(0, weight=1); tree_frame.grid_columnconfigure(0, weight=1)
        self.tree = ttk.Treeview(tree_frame, columns=('trecho', 'nominal', 'admissivel'), show='headings'); self.tree.heading('trecho', text='Trecho'); self.tree.heading('nominal', text='Nominal (A)'); self.tree.heading('admissivel', text='Admissível (A)'); self.tree.pack(fill='both', expand=True)
        botoes_frame = tk.Frame(main_frame); botoes_frame.grid(row=0, column=1, sticky="ns"); ttk.Button(botoes_frame, text="Adicionar Cabo", command=self.adicionar_cabo).pack(pady=5, fill='x'); ttk.Button(botoes_frame, text="Adicionar Múltiplos", command=self.adicionar_multiplos_cabos).pack(pady=5, fill='x'); ttk.Button(botoes_frame, text="Editar Cabo", command=self.editar_cabo).pack(pady=(15, 5), fill='x'); ttk.Button(botoes_frame, text="Remover Cabo", command=self.remover_cabo).pack(pady=5, fill='x')
    def carregar_cabos(self):
        self.tree.delete(*self.tree.get_children()); cabos = self.controller.get_todos_cabos()
        for cabo in cabos: self.tree.insert('', 'end', iid=cabo['id'], values=(cabo['trecho'], cabo['nominal'], cabo['admissivel']))
    def adicionar_cabo(self):
        dialog = CaboDialog(self, title="Adicionar Cabo")
        if dialog.result: trecho, nominal, admissivel = dialog.result; self.controller.adicionar_cabo(trecho, nominal, admissivel); self.carregar_cabos()
    def adicionar_multiplos_cabos(self):
        dialog = MultiCaboDialog(self, title="Adicionar Múltiplos Cabos")
        if dialog.result:
            count = 0
            for trecho, nominal, admissivel in dialog.result: self.controller.adicionar_cabo(trecho, nominal, admissivel); count += 1
            self.carregar_cabos(); messagebox.showinfo("Sucesso", f"{count} cabos adicionados com sucesso.", parent=self)
    def editar_cabo(self):
        selected_id = self.tree.focus()
        if not selected_id: messagebox.showwarning("Atenção", "Selecione um cabo para editar.", parent=self); return
        item = self.tree.item(selected_id, 'values'); dialog = CaboDialog(self, title="Editar Cabo", initial_data=item)
        if dialog.result: trecho, nominal, admissivel = dialog.result; self.controller.editar_cabo(int(selected_id), trecho, nominal, admissivel); self.carregar_cabos()
    def remover_cabo(self):
        selected_id = self.tree.focus()
        if not selected_id: messagebox.showwarning("Atenção", "Selecione um cabo para remover.", parent=self); return
        if messagebox.askyesno("Confirmar", "Deseja remover o cabo selecionado?", parent=self): self.controller.deletar_cabo(int(selected_id)); self.carregar_cabos()

class TabMapeamento(tk.Frame):
    def __init__(self, parent, controller):
        super().__init__(parent); self.controller = controller; self.map_id_map = {}; self.criar_widgets(); self.carregar_mapas()
    def criar_widgets(self):
        main_pane = ttk.PanedWindow(self, orient=tk.HORIZONTAL); main_pane.pack(fill="both", expand=True, padx=10, pady=10)
        map_list_frame = tk.LabelFrame(main_pane, text="Mapas", width=250); main_pane.add(map_list_frame, weight=1); map_list_frame.grid_rowconfigure(0, weight=1); map_list_frame.grid_columnconfigure(0, weight=1)
        self.map_listbox = tk.Listbox(map_list_frame, exportselection=False); self.map_listbox.grid(row=0, column=0, columnspan=2, sticky='nsew'); self.map_listbox.bind('<<ListboxSelect>>', self.on_map_select)
        map_buttons_frame = tk.Frame(map_list_frame); map_buttons_frame.grid(row=1, column=0, columnspan=2, sticky='ew'); ttk.Button(map_buttons_frame, text="Adicionar", command=self.adicionar_mapa).pack(side='left', expand=True, fill='x'); ttk.Button(map_buttons_frame, text="Renomear", command=self.renomear_mapa).pack(side='left', expand=True, fill='x'); ttk.Button(map_buttons_frame, text="Remover", command=self.remover_mapa).pack(side='left', expand=True, fill='x'); ttk.Button(map_buttons_frame, text="Exportar", command=self.exportar_mapeamento_para_excel).pack(side='left', expand=True, fill='x')
        tree_frame = tk.LabelFrame(main_pane, text="Mapeamento de Campos para o Excel"); main_pane.add(tree_frame, weight=3); tree_frame.grid_columnconfigure(0, weight=1); tree_frame.grid_rowconfigure(0, weight=1)
        self.tree = ttk.Treeview(tree_frame, columns=('campo', 'aba', 'celula'), show='headings'); self.tree.heading('campo', text='Campo no Aplicativo'); self.tree.heading('aba', text='Aba no Excel'); self.tree.heading('celula', text='Célula no Excel'); self.tree.column('campo', width=250); self.tree.column('aba', width=150); self.tree.column('celula', width=100); self.tree.grid(row=0, column=0, sticky='nsew'); self.tree.bind("<Double-1>", self.on_edit_mapping)
        botoes_frame = tk.Frame(tree_frame); botoes_frame.grid(row=1, column=0, sticky="ew"); ttk.Button(botoes_frame, text="Adicionar Múltiplos", command=self.adicionar_mapeamentos).pack(side='left', padx=5, pady=5); ttk.Button(botoes_frame, text="Remover Selecionado", command=self.remover_mapeamento).pack(side='left', padx=5, pady=5); ttk.Button(botoes_frame, text="Remover Tudo", command=self.remover_todos_mapeamentos).pack(side='left', padx=5, pady=5)
    def carregar_mapas(self):
        self.map_listbox.delete(0, tk.END); self.map_id_map.clear(); mapas = self.controller.get_all_map_names()
        if not mapas: return
        for mapa in mapas: idx = self.map_listbox.size(); self.map_listbox.insert(tk.END, mapa['nome']); self.map_id_map[idx] = mapa['id']
        if self.map_listbox.size() > 0: self.map_listbox.selection_set(0); self.on_map_select(None)
    def on_map_select(self, event):
        sel = self.map_listbox.curselection()
        if not sel: self.tree.delete(*self.tree.get_children()); return
        map_id = self.map_id_map.get(sel[0])
        if map_id: self.carregar_mapeamento(map_id)
    def adicionar_mapa(self):
        nome = simpledialog.askstring("Novo Mapa", "Digite o nome para o novo mapa:", parent=self)
        if nome and nome.strip(): self.controller.add_map_name(nome.strip()); self.carregar_mapas()
    def renomear_mapa(self):
        sel = self.map_listbox.curselection();
        if not sel: return
        map_id = self.map_id_map.get(sel[0]); nome_atual = self.map_listbox.get(sel[0]); novo_nome = simpledialog.askstring("Renomear Mapa", "Novo nome:", initialvalue=nome_atual, parent=self)
        if novo_nome and novo_nome.strip(): self.controller.rename_map_name(map_id, novo_nome.strip()); self.carregar_mapas()
    def remover_mapa(self):
        sel = self.map_listbox.curselection();
        if not sel: return
        map_id = self.map_id_map.get(sel[0]); nome = self.map_listbox.get(sel[0])
        if messagebox.askyesno("Confirmar", f"Deseja remover o mapa '{nome}' e todos os seus mapeamentos?", icon='warning', parent=self): self.controller.delete_map(map_id); self.carregar_mapas()
    def carregar_mapeamento(self, map_id):
        self.tree.delete(*self.tree.get_children()); mapeamentos = self.controller.get_mappings_for_map(map_id)
        if not mapeamentos: return
        for item in mapeamentos: self.tree.insert('', 'end', iid=item['id'], values=(item['campo_app'], item['aba_excel'], item['celula_excel']))
    def adicionar_mapeamentos(self):
        sel = self.map_listbox.curselection()
        if not sel: messagebox.showwarning("Atenção", "Selecione um mapa à esquerda.", parent=self); return
        map_id = self.map_id_map.get(sel[0]); dialog = MultiMapDialog(self, title="Adicionar Mapeamentos")
        if dialog.result: self.controller.add_mappings_to_map(map_id, dialog.result); self.carregar_mapeamento(map_id); messagebox.showinfo("Sucesso", f"{len(dialog.result)} mapeamentos foram salvos.", parent=self)
    def remover_mapeamento(self):
        selected_id = self.tree.focus()
        if not selected_id: messagebox.showwarning("Atenção", "Selecione um mapeamento para remover.", parent=self); return
        if messagebox.askyesno("Confirmar", "Deseja remover o mapeamento selecionado?", parent=self): self.controller.deletar_mapeamento(int(selected_id)); self.on_map_select(None)
    def remover_todos_mapeamentos(self):
        sel = self.map_listbox.curselection()
        if not sel: messagebox.showwarning("Atenção", "Selecione um mapa para limpar.", parent=self); return
        map_id = self.map_id_map.get(sel[0])
        if messagebox.askyesno("Confirmar Remoção Total", "Tem certeza que deseja remover TODOS os mapeamentos deste mapa?", icon='warning', parent=self): self.controller.deletar_mapeamentos_por_mapa(map_id); self.on_map_select(None)
    def on_edit_mapping(self, event):
        selected_id = self.tree.focus()
        if not selected_id: return
        item_values = self.tree.item(selected_id, 'values'); initial_data = {'campo': item_values[0], 'aba': item_values[1], 'celula': item_values[2]}
        dialog = MapEditDialog(self, "Editar Mapeamento", initial_data=initial_data)
        if dialog.result: self.controller.update_mapping(int(selected_id), dialog.result); self.on_map_select(None)
    def exportar_mapeamento_para_excel(self):
        sel = self.map_listbox.curselection()
        if not sel: messagebox.showwarning("Atenção", "Selecione um mapa para exportar.", parent=self); return
        map_id = self.map_id_map.get(sel[0]); map_name = self.map_listbox.get(sel[0]); mapeamentos = self.controller.get_mappings_for_map(map_id)
        if not mapeamentos: messagebox.showinfo("Vazio", "Este mapa não possui mapeamentos para exportar.", parent=self); return
        filepath = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel Workbook", "*.xlsx")], title="Salvar Mapeamento Como...", initialfile=f"Mapeamento_{map_name}.xlsx")
        if not filepath: return
        try:
            wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Mapeamento"; ws.append(['Campo no Aplicativo', 'Aba no Excel', 'Célula no Excel'])
            for item in mapeamentos: ws.append([item['campo_app'], item['aba_excel'], item['celula_excel']])
            wb.save(filepath); messagebox.showinfo("Sucesso", f"Mapeamento exportado com sucesso para:\n{filepath}", parent=self)
        except Exception as e: messagebox.showerror("Erro na Exportação", f"Ocorreu um erro ao salvar o arquivo:\n{e}", parent=self)
